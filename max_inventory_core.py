import os
import sys
import json
import math
import subprocess
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Formatting utilities
def safe_set_cell(ws, r, c, val=None):
    if (r, c) in ws._cells and type(ws._cells[(r, c)]).__name__ == 'MergedCell':
        del ws._cells[(r, c)]
    cell = ws.cell(row=r, column=c)
    if val is not None:
        cell.value = val
    return cell

def roundup_price(val, to_nearest=1000):
    if not val or val <= 0:
        return 0
    return int(math.ceil(val / to_nearest) * to_nearest)

def calculate_prices(cost, margin_wa=0.35, margin_ml=0.35, comm_ml=0.16):
    cost = float(cost or 0)
    margin_wa = float(margin_wa or 0.35)
    margin_ml = float(margin_ml or 0.35)
    comm_ml = float(comm_ml or 0.16)

    # WhatsApp / Direct price
    if margin_wa < 1:
        raw_direct = cost / (1.0 - margin_wa)
    else:
        raw_direct = cost
    price_direct = roundup_price(raw_direct, 1000)

    # MercadoLibre price
    shipping = 20000 if price_direct >= 90000 else 0
    divisor_ml = 1.0 - margin_ml - comm_ml
    if divisor_ml > 0:
        raw_ml = (cost + shipping) / divisor_ml
    else:
        raw_ml = price_direct
    price_ml = roundup_price(raw_ml, 1000)

    # Facebook price
    price_fb = price_direct

    # Direct Profit & Real Margin
    profit_direct = price_direct - cost
    real_margin = (profit_direct / price_direct) if price_direct > 0 else 0.0

    return {
        'precio_directo_whatsapp': price_direct,
        'precio_mercadolibre': price_ml,
        'precio_facebook': price_fb,
        'utilidad_directa': profit_direct,
        'margen_real_directo': real_margin
    }

def generate_wa_message(item):
    name = (item.get('name') or '').strip()
    brand = (item.get('brand') or '').strip()
    specs = (item.get('specs_amigables') or '').strip()
    estado = (item.get('estado') or 'Nuevo').strip()
    garantia = (item.get('garantia') or 'Garantía oficial').strip()
    barcode = (item.get('barcode') or '').strip()
    part = (item.get('part_number') or '').strip()
    
    price_dir = item.get('precio_directo_whatsapp', 0)
    price_ml = item.get('precio_mercadolibre', 0)
    
    format_cop = lambda v: f"${int(v):,}".replace(',', '.')
    
    title_upper = f"*{name.upper()} – COTIZACIÓN*"
    
    # Format specs
    specs_lines = []
    if specs and specs != 'N/A':
        raw_bullets = [s.strip() for s in specs.replace('•', '\n').split('\n') if s.strip()]
        for b in raw_bullets[:4]:
            if not b.startswith('•'):
                specs_lines.append(f"• {b}")
            else:
                specs_lines.append(b)
    
    if not specs_lines:
        specs_lines = [f"• Marca: {brand}", "• Producto 100% garantizado y probado en taller."]
        
    specs_text = "\n".join(specs_lines)
    
    ref_line = ""
    if barcode and barcode != 'N/A' and barcode != 'Sin código':
        ref_line = f"\n• *Código / EAN:* {barcode}"
    elif part and part != 'N/A':
        ref_line = f"\n• *Referencia / Modelo:* {part}"
        
    estado_text = f"• *Estado:* {estado}"
    if 'reman' in estado.lower():
        estado_text = f"• *Estado:* Equipo remanufacturado en perfecto estado estético, 100% probado en mostrador."
    elif 'nuevo' in estado.lower():
        estado_text = f"• *Estado:* Producto nuevo, sellado en caja."
        
    msg = (
        f"¡Hola! Con gusto te paso los datos del equipo:\n\n"
        f"🏷️ {title_upper}\n\n"
        f"{specs_text}\n\n"
        f"{estado_text}\n"
        f"• *Garantía:* {garantia}{ref_line}\n\n"
        f"💰 *Precio en efectivo o transferencia:* {format_cop(price_dir)} COP\n"
        f"💳 *(Opción con tarjeta a cuotas o por MercadoLibre:* {format_cop(price_ml)} COP)\n\n"
        f"¿Te gustaría que te lo separe o tienes alguna pregunta sobre el equipo? ¡Con gusto te ayudo!"
    )
    return msg

def load_products_json(json_path):
    if os.path.exists(json_path):
        with open(json_path, 'r', encoding='utf-8') as f:
            return json.load(f)
    return []

def save_products_json(products, json_path):
    # Recalculate prices and wa messages
    for p in products:
        cost = p.get('costo_total_unitario', 0)
        m_wa = p.get('margen_whatsapp', 0.35)
        m_ml = p.get('margen_mercadolibre', 0.35)
        c_ml = p.get('comision_mercadolibre', 0.16)
        
        calc = calculate_prices(cost, m_wa, m_ml, c_ml)
        p.update(calc)
        p['wa_samuel_message'] = generate_wa_message(p)
        p['disponibilidad'] = '🟢 En Stock' if (p.get('stock_actual', 0) > 0) else '⚪ Agotado'

    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(products, f, ensure_ascii=False, indent=2)
    return products

def sync_json_to_excel(products, excel_path):
    wb = openpyxl.load_workbook(excel_path)
    ws = wb['📦 Inventario Maestro']
    
    # Clear existing data rows starting from row 4
    max_r = max(ws.max_row, 100)
    for r in range(4, max_r + 1):
        for c in range(1, 25):
            safe_set_cell(ws, r, c).value = None

    # Styles
    font_body = Font(name='Segoe UI', size=9)
    font_bold = Font(name='Segoe UI', size=9, bold=True)
    font_sku = Font(name='Segoe UI Semibold', size=9, bold=True, color='D91F26')
    
    border_thin = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )
    
    align_left = Alignment(horizontal='left', vertical='center')
    align_center = Alignment(horizontal='center', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')
    
    for idx, p in enumerate(products, start=4):
        r = idx
        sku = p.get('sku') or f"MTS-{r-3:04d}"
        name = p.get('name', '')
        cat = p.get('category', '')
        brand = p.get('brand', '')
        part = p.get('part_number', '')
        color = p.get('color', '')
        barcode = p.get('barcode', '')
        estado = p.get('estado', 'NUEVO')
        stock = p.get('stock_actual', 0)
        disp = '🟢 En Stock' if stock > 0 else '⚪ Agotado'
        hist_comp = p.get('historico_comprado', stock)
        ubic = p.get('ubicacion', 'Taller')
        costo = p.get('costo_total_unitario', 0)
        m_wa = p.get('margen_whatsapp', 0.35)
        com_ml = p.get('comision_mercadolibre', 0.16)
        garantia = p.get('garantia', 'Garantía oficial')
        img = p.get('img_path', '')
        specs = p.get('specs_amigables', '')
        notas = p.get('notas', '')

        # Set values
        safe_set_cell(ws, r, 1, sku).font = font_sku
        safe_set_cell(ws, r, 2, name).font = font_bold
        safe_set_cell(ws, r, 3, cat).font = font_body
        safe_set_cell(ws, r, 4, brand).font = font_body
        safe_set_cell(ws, r, 5, part).font = font_body
        safe_set_cell(ws, r, 6, color).font = font_body
        safe_set_cell(ws, r, 7, barcode).font = font_body
        safe_set_cell(ws, r, 8, estado).font = font_body
        safe_set_cell(ws, r, 9, disp).font = font_bold
        safe_set_cell(ws, r, 10, stock).font = font_bold
        safe_set_cell(ws, r, 11, hist_comp).font = font_body
        safe_set_cell(ws, r, 12, ubic).font = font_body
        safe_set_cell(ws, r, 13, costo).font = font_body
        safe_set_cell(ws, r, 14, m_wa).font = font_body
        
        # Formulas for prices
        safe_set_cell(ws, r, 15, f"=IF(N{r}<1, ROUNDUP(M{r}/(1-N{r}), -3), M{r})").font = font_bold
        safe_set_cell(ws, r, 16, com_ml).font = font_body
        safe_set_cell(ws, r, 17, f"=IF((1-N{r}-P{r})>0, ROUNDUP((M{r}+IF(O{r}>=90000,20000,0))/(1-N{r}-P{r}), -3), O{r})").font = font_body
        safe_set_cell(ws, r, 18, f"=O{r}").font = font_body
        safe_set_cell(ws, r, 19, f"=O{r}-M{r}").font = font_body
        safe_set_cell(ws, r, 20, f"=IF(O{r}>0, (O{r}-M{r})/O{r}, 0)").font = font_body
        
        safe_set_cell(ws, r, 21, garantia).font = font_body
        safe_set_cell(ws, r, 22, img).font = font_body
        safe_set_cell(ws, r, 23, specs).font = font_body
        safe_set_cell(ws, r, 24, notas).font = font_body

        # Alignments & Number formats
        for c in range(1, 25):
            cell = safe_set_cell(ws, r, c)
            cell.border = border_thin
            if c in [1, 5, 6, 7, 8, 9, 10, 11, 12, 21]:
                cell.alignment = align_center
            elif c in [13, 15, 17, 18, 19]:
                cell.alignment = align_right
                cell.number_format = '$#,##0'
            elif c in [14, 16, 20]:
                cell.alignment = align_right
                cell.number_format = '0.0%'
            else:
                cell.alignment = align_left

    wb.save(excel_path)
    return True

def sync_excel_to_json(excel_path, json_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['📦 Inventario Maestro']
    
    products = []
    for r in range(4, ws.max_row + 1):
        sku = safe_set_cell(ws, r, 1).value
        name = safe_set_cell(ws, r, 2).value
        if not sku or not name:
            continue
            
        cat = safe_set_cell(ws, r, 3).value or 'General'
        brand = safe_set_cell(ws, r, 4).value or 'Genérica'
        part = safe_set_cell(ws, r, 5).value or ''
        color = safe_set_cell(ws, r, 6).value or ''
        barcode = safe_set_cell(ws, r, 7).value or ''
        estado = safe_set_cell(ws, r, 8).value or 'NUEVO'
        stock = int(safe_set_cell(ws, r, 10).value or 0)
        hist_comp = int(safe_set_cell(ws, r, 11).value or stock)
        ubic = safe_set_cell(ws, r, 12).value or 'Taller'
        costo = float(safe_set_cell(ws, r, 13).value or 0)
        m_wa = float(safe_set_cell(ws, r, 14).value or 0.35)
        com_ml = float(safe_set_cell(ws, r, 16).value or 0.16)
        garantia = safe_set_cell(ws, r, 21).value or 'Garantía oficial'
        img = safe_set_cell(ws, r, 22).value or ''
        specs = safe_set_cell(ws, r, 23).value or ''
        notas = safe_set_cell(ws, r, 24).value or ''

        calc = calculate_prices(costo, m_wa, m_wa, com_ml)
        
        item = {
            'sku': str(sku).strip(),
            'name': str(name).strip(),
            'category': str(cat).strip(),
            'brand': str(brand).strip(),
            'part_number': str(part).strip(),
            'color': str(color).strip(),
            'barcode': str(barcode).strip(),
            'estado': str(estado).strip(),
            'disponibilidad': '🟢 En Stock' if stock > 0 else '⚪ Agotado',
            'stock_actual': stock,
            'historico_comprado': hist_comp,
            'ubicacion': str(ubic).strip(),
            'costo_total_unitario': costo,
            'margen_whatsapp': m_wa,
            'margen_mercadolibre': m_wa,
            'comision_mercadolibre': com_ml,
            'garantia': str(garantia).strip(),
            'img_path': str(img).strip(),
            'specs_amigables': str(specs).strip(),
            'notas': str(notas).strip()
        }
        item.update(calc)
        item['wa_samuel_message'] = generate_wa_message(item)
        products.append(item)

    save_products_json(products, json_path)
    return products
